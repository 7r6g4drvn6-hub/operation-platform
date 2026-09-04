from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ARCH = DOCS / "architecture"
OUTPUT = DOCS / "Operation_Platform_V1.1_Architecture.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def add_heading(document, text, level=1):
    heading = document.add_heading(text, level=level)
    heading.style.font.color.rgb = RGBColor(18, 59, 105)
    return heading


def add_code(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.left_indent = Inches(0.18)
    run = paragraph.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(8.5)
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_diagram(document, filename, caption):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(ARCH / filename), width=Inches(6.4))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.runs[0].italic = True
    caption_paragraph.runs[0].font.size = Pt(8.5)
    caption_paragraph.runs[0].font.color.rgb = RGBColor(96, 125, 153)


document = Document()
section = document.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.68)
section.right_margin = Inches(0.68)

styles = document.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor(23, 43, 77)
styles["Title"].font.name = "Arial"
styles["Title"].font.size = Pt(28)
styles["Title"].font.color.rgb = RGBColor(18, 59, 105)

title = document.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Operation Platform V1.1")
subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_run = subtitle.add_run("China Operation Intelligence Platform\nAPI Architecture, Signal Data Flow, and Overview Architecture")
subtitle_run.bold = True
subtitle_run.font.size = Pt(15)
subtitle_run.font.color.rgb = RGBColor(47, 128, 183)
meta = document.add_paragraph("Evidence-driven Operation Platform · Mock Prototype Specification\nGenerated 2026-09-03")
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.runs[0].font.color.rgb = RGBColor(92, 113, 136)

document.add_paragraph()
callout = document.add_table(rows=1, cols=1)
callout.autofit = True
set_cell_shading(callout.cell(0, 0), "EEF7FB")
callout.cell(0, 0).paragraphs[0].add_run(
    "Product position: Operation Platform is not an ITSM, Incident Management, Monitoring, or Probe Management system. "
    "It connects evidence to accountable action: Detect → Understand → Act → Prove → Improve."
).bold = True

add_heading(document, "1. Executive Architecture")
document.add_paragraph(
    "V1.1 adds a Signal Intelligence Layer upstream of the existing V1 P0 closed loop. External China-region events "
    "are normalized into a common Signal object, correlated into Problem Candidates, and then passed into the existing "
    "Problem → Operation → Action → Evidence → Verification → Outcome → History workflow."
)
add_code(document, "External Sources → API / Integration → Source Adapter → Signal Service → Correlation → Problem → Operation → Action → Evidence → Verification → Outcome → China Operation Overview → History")

document.add_page_break()
add_heading(document, "2. API Architecture")
add_diagram(document, "Operation_Platform_V1.1_API_Architecture.png", "Figure 1. External sources, API boundary, adapters, Signal Intelligence, Operation core, and storage.")
add_heading(document, "Reserved API contracts", level=2)
add_code(document, """POST /api/signals
GET  /api/signals
GET  /api/signals/:id
POST /api/problems/from-signals
POST /api/integration/probe/result
POST /api/integration/incident
POST /api/integration/smo-ticket
POST /api/integration/availability
POST /api/integration/release""")
document.add_paragraph("These are mock contracts only. No real external API, Kafka, database, email ingestion, or AI correlation is introduced in V1.1.")

document.add_page_break()
add_heading(document, "3. Signal Data Flow")
add_diagram(document, "Operation_Platform_V1.1_Data_Flow.png", "Figure 2. SIG-0001 is correlated with independent context, then moves through action, verification, and Outcome.")
add_heading(document, "Reference demo flow", level=2)
add_code(document, """SIG-0001 (Probe, 12.8% failure, China)
→ Correlation (same service / region / time window)
→ PRB-204 Approval Download Instability
→ OP-102 Improve Approval Download Stability
→ Action + Probe / Availability / Release Evidence
→ Verification Probe (expected < 1%, actual 0.6%)
→ OUT-001 Objective Achieved
→ Operation Completed + Problem Resolved + History""")

document.add_page_break()
add_heading(document, "4. China Operation Overview")
add_diagram(document, "Operation_Platform_V1.1_Overview_Architecture.png", "Figure 3. Overview aggregates Health, Business Impact, Outcome, Trend, and Closed-loop performance.")
document.add_paragraph(
    "The Overview answers: How healthy is China operation? What did we resolve? Did our Operations actually improve "
    "the situation? Command Center remains a separate attention queue answering what needs action now."
)

add_heading(document, "5. Domain Model")
table = document.add_table(rows=1, cols=3)
table.style = "Table Grid"
headers = ["Object", "Role", "Key relationship"]
for index, header in enumerate(headers):
    cell = table.rows[0].cells[index]
    set_cell_shading(cell, "123B69")
    run = cell.paragraphs[0].add_run(header)
    run.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
rows = [
    ("Signal", "Observed abnormal condition or business/technical event", "Source event → Signal → Problem"),
    ("Incident", "Backend-reported technical context", "Incident → Signal; may contribute to Problem"),
    ("Problem", "Correlated pattern across signals, impact, and evidence", "Problem → Operation"),
    ("Operation", "Accountable effort with objective, owner, actions, and verification", "Operation → Action / Evidence / Outcome"),
    ("Action", "Executable task with expected result and owner", "Action → Evidence"),
    ("Probe", "Independent evidence generator", "Probe Result → Signal and Evidence"),
    ("Evidence", "Objective proof of condition, action, or outcome", "Evidence → Verification"),
    ("Outcome", "Measured conclusion after verification", "Outcome → History"),
]
for values in rows:
    cells = table.add_row().cells
    for index, value in enumerate(values):
        cells[index].text = value

document.add_page_break()
add_heading(document, "Signal contract", level=2)
add_code(document, """{
  id, source, type, service, region, environment,
  severity, status,
  metric: { name, value, threshold, unit },
  impact: { userAffected, businessImpact },
  timestamp, relatedObjects, problemId, correlationScore
}""")

add_heading(document, "6. Rules and Boundaries")
add_bullets(document, [
    "Signal is the unified entry object; it is not an Incident and not Evidence.",
    "Rule-based correlation uses shared service, region, and time window; no AI is required.",
    "A Problem can exist without a Backend Incident. Unreported Problems are first-class attention items.",
    "Probe remains an independent product. Operation consumes Probe Results as Evidence.",
    "Action completion alone never closes an Operation. Only passed verification can create an Outcome.",
    "All Signal → Problem → Operation → Outcome state transitions are recorded in historyEvents.",
])

add_heading(document, "7. Prototype Scope and Limitations")
document.add_paragraph(
    "V1.1 uses Mock Data and a browser-local shared State Layer. The API files define the replacement boundary for "
    "future real services while the UI remains decoupled from source-specific payloads."
)
add_bullets(document, [
    "No real Probe, SMO, Incident, Availability, or Release API connections.",
    "No production database, Kafka, email parser, AI root cause, predictive detection, or automatic Probe trigger.",
    "server.js remains unchanged.",
])

document.save(OUTPUT)
print(OUTPUT)
