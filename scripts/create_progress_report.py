from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Operation_Platform_Product_Progress_Report.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def add_run(paragraph, text, bold=False, color=None, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return run


def set_style_fonts(style, latin="Aptos", cjk="Hiragino Sans GB"):
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)
    rfonts.set(qn("w:cs"), cjk)


def set_document_run_fonts(doc, cjk="Hiragino Sans GB"):
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    for paragraph in paragraphs:
        latin = "Menlo" if paragraph.style and paragraph.style.name == "Code Block" else "Aptos"
        for run in paragraph.runs:
            rpr = run._r.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.insert(0, rfonts)
            rfonts.set(qn("w:ascii"), latin)
            rfonts.set(qn("w:hAnsi"), latin)
            rfonts.set(qn("w:eastAsia"), cjk)
            rfonts.set(qn("w:cs"), cjk)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_code(doc, text, font_size=8.5):
    p = doc.add_paragraph()
    p.style = doc.styles["Code Block"]
    p.paragraph_format.space_after = Pt(8)
    p.add_run(text)
    for run in p.runs:
        run.font.size = Pt(font_size)
    return p


def add_figure(doc, path, caption, width=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    r = c.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(89, 89, 89)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_cant_split(hdr)
    for idx, head in enumerate(headers):
        cell = hdr.cells[idx]
        shade(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            cell.width = Inches(widths[idx])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(head)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)
    for ridx, row in enumerate(rows):
        added_row = table.add_row()
        set_cant_split(added_row)
        cells = added_row.cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[idx])
            if ridx % 2 == 1:
                shade(cell, "F3F6F9")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    set_keep_with_next(p)
    return p


def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        add_run(p, bold_lead, bold=True)
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(32, 32, 32)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    set_style_fonts(normal)
    for style_name, size, color in (("Title", 24, "000000"), ("Heading 1", 16, "000000"), ("Heading 2", 12.5, "000000"), ("Heading 3", 11, "000000")):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        set_style_fonts(style)
    if "Code Block" not in [s.name for s in doc.styles]:
        code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = doc.styles["Code Block"]
    code.font.name = "Courier New"
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor(45, 45, 45)
    set_style_fonts(code, latin="Menlo")

    # Header and footer
    header = section.header.paragraphs[0]
    header.text = "OPERATION PLATFORM  |  PRODUCT PROGRESS REPORT"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(102, 102, 102)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Operation Platform V1.2  ·  2026-09-04")
    for r in footer.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(102, 102, 102)

    title = doc.add_paragraph(style="Title")
    title.add_run("Operation Platform 产品进度与交付路线汇报")
    title.paragraph_format.space_after = Pt(8)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(14)
    add_run(sub, "项目版本：V1.2 Prototype 及真实闭环交付准备", bold=True, color="1F4E78", size=11)
    para(doc, "本报告向产品、Operation、Backend、Frontend、Probe 和管理团队说明当前已经交付的产品能力、仍然存在的生产化差距，以及从数据库建设到第一条真实闭环的下一步执行路线。报告明确区分可演示的 Prototype 与尚未完成的生产系统。")

    heading(doc, "一、汇报结论", 1)
    para(doc, "当前产品已经完成从 Incident 管理页面向 Evidence driven Operation Platform 原型的产品骨架升级，并已发布到 GitHub Pages，任何人都可以在线查看。原型能够演示 Detect、Understand、Act、Verify、Outcome 的基本闭环。")
    add_code(doc, "Detect → Understand → Act → Verify → Outcome", 12)
    para(doc, "当前完成的是 UI、信息架构、领域模型、Mock Data 和交互演示，不是生产系统。生产化尚未开始，核心缺口是 PostgreSQL、正式 API、Backend 状态机、鉴权、真实 Probe 回调、服务端审计和前端 API 迁移。")
    add_code(doc, "数据库 → API → Backend → Probe 接入 → 前端改造 → 第一条真实闭环 → 验收", 11)

    heading(doc, "二、当前交付状态", 1)
    add_table(doc, ["能力", "当前状态", "已完成内容", "生产化缺口"], [
        ["产品模型", "已完成原型", "Signal、Incident、Problem、Operation、Action、Probe、Evidence、Verification、Outcome、History", "落到服务端实体和约束"],
        ["Command Center", "已完成原型", "Critical/High Risk、Unreported Problems、Active Operations、Verification attention", "仍读取浏览器 Mock State"],
        ["Problem Detail", "已完成原型", "Impact、Signals、Incidents、Probe、Availability、Recommendation、Active Operations、Timeline", "推荐逻辑需要服务端规则和真实 Evidence"],
        ["Operation Detail", "已完成原型", "Objective、Owner、Actions、Progress、Evidence、Verification、Outcome", "需要后端状态机和权限控制"],
        ["Unreported Problem", "已完成原型", "Probe 异常且无 Backend Incident 的问题可单独展示和推进", "需要真实 Signal ingestion 和 Problem API"],
        ["Probe 边界", "已定义并演示", "Probe 保持独立，作为 Evidence Generator", "尚无真实出站请求和回调"],
        ["Closed loop", "可演示", "Failed Verification 不关闭；Passed Verification 产生 Outcome", "尚无事务一致性和数据库审计"],
        ["Mock Data", "已完成", "China / MIB3 Approval、Problems、Incidents、Probe Results、Operations、Actions、Evidence", "不能作为生产数据源"],
        ["文档和架构图", "已完成", "V1.1 API、Data Flow、Overview 图及交付计划", "转化为 OpenAPI、ERD 和运行手册"],
        ["GitHub 发布", "已完成", "Public repository 和 GitHub Pages 在线预览", "GitHub Pages 只适合静态 Prototype"],
    ], widths=[1.05, 0.95, 2.6, 1.75])

    heading(doc, "三、已经完成了什么", 1)
    heading(doc, "3.1 产品定位和核心原则", 2)
    para(doc, "产品已经明确不是另一个 Incident Management Dashboard，而是以 Evidence 为核心的运营闭环平台：")
    add_code(doc, "Signal → Problem → Operation → Action → Evidence → Verification → Outcome", 10)
    add_bullets(doc, [
        "Incident 是 Signal Source，不是唯一入口。",
        "Problem 可以在 Incident 数量为 0 时存在。",
        "Probe 是独立产品，不复制 Probe 执行逻辑。",
        "Action 完成不等于 Operation 完成。",
        "只有通过明确 KPI 验证的 Verification 才能创建 Outcome。",
        "Overview 解释结果和业务影响，Command Center 负责当前注意力和风险排序。",
    ])
    heading(doc, "3.2 当前前端原型", 2)
    para(doc, "当前页面和路由覆盖 China Operation Overview、Command Center、Signals、Problem List / Detail、Operations / Detail、Evidence、Incidents、Probe、Outcome Detail 和 Closed-loop History。")
    add_bullets(doc, [
        "Problem Detail 中的 Recommended Next Action。",
        "Problem Detail 中的 Active Operations、负责人、状态、进度和 Actions。",
        "Command Center 中的一等 Unreported Problems 区域。",
        "Signal → Problem → Operation → Action → Probe → Evidence → Verification → Outcome 导航。",
        "Verification Failed 后保持 Operation 在 Verifying，并建议 Corrective Action。",
        "Verification Passed 后显示 Outcome、Before/After 和业务影响。",
    ])
    heading(doc, "3.3 真实场景 Mock Data", 2)
    add_code(doc, "Service: MIB3 Approval\nRegion: China\nJourney: Approval Download\nInitial Probe failure rate: 12.8%\nInitial Availability: 96.1%\nTarget failure rate: < 1%\nTarget Availability: > 99.7%\nInitial Incident count: 0", 9.5)
    para(doc, "该场景明确展示了“Probe 已经发现问题，但 Backend 尚未创建 Incident”的差异化能力。")
    heading(doc, "3.4 GitHub 交付", 2)
    add_bullets(doc, [
        "代码仓库：https://github.com/7r6g4drvn6-hub/operation-platform",
        "在线 Prototype：https://7r6g4drvn6-hub.github.io/operation-platform/",
        "分支：main；本地运行：node server.js",
    ])
    para(doc, "GitHub Pages 当前发布的是静态 UI Prototype。每个访问者使用自己的浏览器 Mock State，数据不会在访问者之间同步，也没有真实 API、数据库或生产鉴权。")

    heading(doc, "四、当前实现架构", 1)
    para(doc, "当前架构用于快速验证产品模型和用户流程，业务状态保存在浏览器 Prototype State 中。优点是启动快、适合演示；缺点是没有跨用户一致性、服务端权限、可靠事件接入和审计能力。")
    add_code(doc, "访问者浏览器\n      ↓\nGitHub Pages 或本地 server.js\n      ├── index.html\n      ├── styles.css / styles-v11.css / styles-v12.css\n      ├── app.js（Hash router 和页面渲染）\n      ├── prototype-state.js（Mock State）\n      └── api/ mock adapters\n              ↓\n       Browser local state\n              ├── Signals\n              ├── Problems\n              ├── Operations / Actions\n              ├── Evidence / Verification\n              └── Outcomes / History", 8.7)
    add_figure(doc, ROOT / "docs/architecture/Operation_Platform_V1.1_Overview_Architecture.png", "图 1  当前 Prototype 与 China Operation Overview 的运行关系")

    heading(doc, "五、目标生产架构", 1)
    para(doc, "第一版生产后端建议采用模块化单体，保持现有领域模型和 Probe 独立边界。PostgreSQL 是 source of truth；API 负责鉴权、校验、幂等和状态命令；前端只消费服务端 DTO。")
    add_code(doc, "┌─────────────────────────────────────────────────────────────┐\n│ External Signal Sources                                     │\n│ Probe · Availability · Incident · SMO · Release · KPI       │\n└──────────────────────────┬──────────────────────────────────┘\n                           │ authenticated events\n                           ▼\n┌─────────────────────────────────────────────────────────────┐\n│ Operation API                                                │\n│ Gateway · Auth · Validation · Idempotency · Trace ID         │\n└──────────────┬──────────────────────────┬───────────────────┘\n               ▼                          ▼\n┌──────────────────────────┐   ┌─────────────────────────────┐\n│ Integration Adapters     │   │ Domain Commands and Queries  │\n│ raw event → normalized   │   │ Signal · Problem · Operation │\n│ event                    │   │ Action · Evidence · Outcome │\n└──────────────┬───────────┘   └──────────────┬──────────────┘\n               ▼                              ▼\n┌─────────────────────────────────────────────────────────────┐\n│ PostgreSQL Source of Truth                                  │\n│ source_events · signals · problems · operations · evidence   │\n│ verification_runs · outcomes · history_events · outbox       │\n└──────────────┬────────────────────────────────────────────────┘\n               ▼\n┌──────────────────────────┐   ┌─────────────────────────────┐\n│ Correlation and Rules    │   │ Probe Integration Boundary   │\n│ service/region/time      │◄──┤ request run · callback result│\n│ window/metric context    │   │ Probe remains independent   │\n└──────────────┬───────────┘   └──────────────┬──────────────┘\n               ▼                              ▼\n┌─────────────────────────────────────────────────────────────┐\n│ Problem → Operation → Action → Evidence → Verification      │\n│                              └──────────────→ Outcome        │\n└──────────────────────────────┬──────────────────────────────┘\n                               ▼\n┌─────────────────────────────────────────────────────────────┐\n│ Read Models and UI                                           │\n│ Command Center · Problem Detail · Operation Detail           │\n│ China Overview · Evidence · Outcome · Closed-loop History   │\n└─────────────────────────────────────────────────────────────┘", 7.4)
    add_figure(doc, ROOT / "docs/architecture/Operation_Platform_V1.1_API_Architecture.png", "图 2  目标 API Architecture：来源适配、Signal Intelligence 和领域 API")
    add_figure(doc, ROOT / "docs/architecture/Operation_Platform_V1.1_Data_Flow.png", "图 3  目标 Data Flow：从 Signal Detection 到 Outcome 和 History")

    heading(doc, "六、详细闭环流程", 1)
    heading(doc, "6.1 主路径", 2)
    add_code(doc, "1. Signal detected\n   Probe、Availability、Incident 或其他来源产生原始事件\n        ↓\n2. Event stored and normalized\n   保存 source event，校验、去重，再生成统一 Signal\n        ↓\n3. Problem correlated\n   按服务、地区、环境、时间窗口和指标上下文形成候选\n        ↓\n4. Problem understood\n   Operation Manager 查看 Impact、Evidence、Incident 和推荐动作\n        ↓\n5. Operation created\n   写入 Objective、Owner、KPI、Expected Outcome 和 Evidence Requirement\n        ↓\n6. Action assigned\n   Backend、Probe Ops 或 Product 承担有明确 Expected Result 的任务\n        ↓\n7. Fix or action completed\n   记录 Actual Result，必要时关联 Action Evidence\n        ↓\n8. Verification Probe executed\n   Probe 独立执行，Operation 只接收 Probe Result\n        ↓\n9. Evidence comparison\n   服务端比较 target snapshot 与 actual snapshot\n        ↓\n10. Outcome decided\n    只有授权人显式 PASS 才能创建 Outcome\n        ↓\n11. Problem resolved\n    Operation 完成、Problem 解决、Signals 归档、History 追加", 8.5)
    heading(doc, "6.2 失败验证路径", 2)
    add_code(doc, "Verification result = 4.8% failure rate\n              ↓\nVerification FAILED\n              ├── Operation remains VERIFYING\n              ├── Problem remains open\n              ├── No Outcome created\n              ├── Corrective Action becomes available\n              └── New Verification Run can be started", 9.5)
    heading(doc, "6.3 通过验证路径", 2)
    add_code(doc, "Corrective Action completed\n              ↓\nVerification Probe = 0.4% failure rate\nAvailability = 99.8%\n              ↓\nAuthorized reviewer records PASS\n              ↓\nOne transaction updates:\nVerification PASSED · Outcome VERIFIED · Operation COMPLETED\nProblem RESOLVED · Signals ARCHIVED · History appended · Outbox created\n              ↓\nChina Operation Overview refreshed", 9.5)

    heading(doc, "七、当前实现与生产目标的差距", 1)
    add_table(doc, ["差距", "当前情况", "影响", "解决阶段"], [
        ["Source of truth", "浏览器 local state", "刷新、换设备、换用户后不共享", "数据库、Backend"],
        ["API", "Prototype mock adapters", "无稳定契约，前后端无法并行交付", "API"],
        ["Persistence", "无 PostgreSQL", "无法审计、对账、恢复", "数据库"],
        ["State machine", "主要由前端函数驱动", "客户端可能绕过规则", "Backend"],
        ["Authentication", "无生产鉴权", "无法区分角色和集成来源", "API、Backend"],
        ["Idempotency", "未对真实事件持久化去重", "Probe 重试可能重复生成对象", "数据库、Probe"],
        ["Correlation", "Mock rule 结果", "无服务端候选、置信度和配置", "Backend"],
        ["Probe integration", "仅模拟 Run Probe", "无真实 runId、callback、重试、对账", "Probe 接入"],
        ["Evidence calculation", "Mock 数字", "无法证明真实 KPI 改善", "Backend、真实闭环"],
        ["Dashboard aggregation", "读取 Mock collections", "生产统计不能独立硬编码", "Backend Overview"],
        ["Audit", "Prototype History", "不满足 append-only 审计", "数据库、Backend"],
        ["Deployment", "只有静态 Pages", "不能承载 API、数据库和密钥", "Backend 基础设施"],
    ], widths=[1.2, 1.45, 2.25, 1.45])

    heading(doc, "八、下一步交付路线", 1)
    stages = [
        ("Step 1 数据库", "把 Prototype State 替换为 PostgreSQL source of truth。交付八组迁移、China / MIB3 seed、schema、ERD 和完整性测试。门禁是从零迁移/回滚、Unreported Problem 可保存、Probe callback 幂等、Outcome 完整性和 History append-only。"),
        ("Step 2 API", "冻结 OpenAPI，让 Backend、Frontend 和 Probe 并行实现。核心接口覆盖来源接入、Signal、Problem、Recommendation、Operation、Probe Run、Verification Decision、Overview 和 Command Center。门禁是 OpenAPI lint 零错误及角色、状态、幂等和错误码签字。"),
        ("Step 3 Backend", "实现服务端状态机、领域服务、相关性规则、聚合查询和审计。顺序为服务基础和鉴权 → 幂等接入 → Signal/Correlation/Problem → Operation/Action → Evidence/Verification/Outcome → Overview/History/Outbox。"),
        ("Step 4 Probe 接入", "确认 clientRequestId、runId、eventId、上下文、目标指标、callback 鉴权、重试、时钟偏差和 retention；实现出站请求、回调、结果分类、重试和 reconciliation。"),
        ("Step 5 前端改造", "保留当前视觉语言和信息架构，把业务读写从 local state 切换到 API。生产模式不加载 prototype-state.js，不读写业务 localStorage，所有命令显示服务端状态，并具备 Loading、Empty、Error、Unauthorized、Stale、Retry。"),
        ("Step 6 第一条真实闭环", "以 China / MIB3 Approval / Approval Download 为固定场景，执行真实 Probe 异常 12.8% → Unreported Problem → Operation → Diagnostic Evidence → Fix Evidence → 4.8% 失败验证 → Corrective Action → 0.4% 通过验证 → Outcome → Overview 更新。"),
        ("Step 7 最终验收", "管理层仅通过 Overview 和 Command Center 回答健康、风险、未报告问题、执行中的 Operation、已解决 Problem、结果证据和 KPI 改善。"),
    ]
    add_table(doc, ["阶段", "直接交付和完成重点"], stages, widths=[1.55, 4.8])

    heading(doc, "九、第一条真实闭环验收口径", 1)
    heading(doc, "通过条件", 2)
    add_bullets(doc, [
        "Probe 异常来自独立 Probe 环境，不是数据库 seed。",
        "事件通过真实 API 接收并持久化。",
        "Problem 可在 Incident count 为 0 时创建。",
        "Problem Detail 展示 Recommendation、Evidence、Active Operation 和 Timeline。",
        "Failed Verification 不创建 Outcome、不解决 Problem。",
        "Corrective Action 后的 Passed Verification 只创建一个 Outcome。",
        "PASS 事务同时更新 Verification、Outcome、Operation、Problem、Signals、History、Outbox。",
        "Overview 数字可与明细和数据库记录对账。",
        "没有人工修改数据库 status。",
    ])
    heading(doc, "必须留存的证据", 2)
    add_code(doc, "execution-log.md\napi-transcript.json\nentity-ids.json\ndatabase-checks.sql\nscreenshots/\nprobe-result-links.md\nknown-issues.md", 9.5)

    heading(doc, "十、风险、依赖和需要确认的事项", 1)
    add_table(doc, ["项目", "风险或依赖", "需要确认"], [
        ["Probe 合同", "是否支持 context、callback、签名、重试和 result URL", "Probe Owner、Security"],
        ["KPI 定义", "failure rate、availability、单位和时间窗口必须一致", "Product、Probe、Backend"],
        ["权限", "谁能创建 Problem、记录 PASS、取消 Operation", "Product、Security、Operation"],
        ["部署环境", "PostgreSQL、密钥、OIDC、域名和网络白名单", "Backend、IT、Security"],
        ["数据保留", "raw payload、History、Evidence 的保留期限", "Security、Compliance"],
        ["真实窗口", "China Production 的 Probe 执行和回滚联系人", "Operation、Backend、Probe"],
    ], widths=[1.2, 3.4, 1.75])

    heading(doc, "十一、项目文件地图", 1)
    add_code(doc, "operation-platform/\n├── index.html                         # 页面壳和脚本入口\n├── app.js                             # Hash router 和页面渲染\n├── prototype-state.js                 # 当前 Prototype Mock State\n├── api/                               # 当前 Signal mock adapters\n├── styles.css                         # 现有视觉基础\n├── styles-v11.css / styles-v12.css    # 版本增量样式\n├── server.js                          # 本地静态开发服务器\n├── docs/architecture/                 # API、Data Flow、Overview 架构图\n├── docs/Operation_Platform_Execution_Backlog.md\n└── docs/Operation_Platform_Real_Closed_Loop_Delivery_Plan.md", 8.7)

    heading(doc, "十二、最终判断", 1)
    para(doc, "项目已经完成产品方向验证和可公开查看的 UI Prototype，下一阶段不是继续堆页面，而是把已有模型变成可审计、可重放、可验证的服务端闭环。")
    para(doc, "真正的生产完成标准不是页面数量，而是下面这条链路可以被 API、数据库和 UI 同时证明：")
    add_code(doc, "Detect → Understand → Act → Verify → Outcome", 12)
    para(doc, "只要仍然存在浏览器 local state、手工改状态、Probe 结果无法对账或失败 Verification 可以误关闭 Operation，就不能称为 Operation Platform 生产闭环已经完成。")

    set_document_run_fonts(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
